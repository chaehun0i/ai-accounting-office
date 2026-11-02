"""A CLOSED snapshot can be requested once and read only within its company."""

from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.storage.domain.contracts import StorageUnavailable
from tests.integration.test_financial_statements import closed_company


def test_generation_pins_closed_source_and_replays_without_duplicate(api_client, auth_service):
    owner, writer, period, _, _, _ = closed_company(api_client, auth_service)
    generated = api_client.post(
        "/financial-statements/generate",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json={"period_id": period},
    )
    assert generated.status_code == 201, generated.text
    statement_set = generated.json()
    selected = api_client.get("/accounting/periods", headers=writer)
    assert selected.status_code == 200, selected.text
    period_row = next(row for row in selected.json() if row["id"] == period)
    payload = {
        "report_code": "BALANCE_SHEET",
        "format": "CSV",
        "period_id": period,
        "period_preset": "CUSTOM",
        "as_of": period_row["end_date"],
    }
    key = str(uuid4())
    headers = {**writer, "Idempotency-Key": key}
    first = api_client.post("/reports/generations", headers=headers, json=payload)
    assert first.status_code == 201, first.text
    run = first.json()
    assert run["statement_set_id"] == statement_set["id"]
    assert run["trial_balance_digest"] == statement_set["trial_balance_digest"]
    assert run["period_mode"] == "AS_OF"
    assert run["status"] == "COMPLETED"
    artifact = api_client.get(f"/reports/generations/{run['id']}/artifact", headers=writer)
    assert artifact.status_code == 200
    assert artifact.json()["file_name"].endswith(".csv")
    assert "storage_key" not in artifact.json()
    downloaded = api_client.get(f"/reports/generations/{run['id']}/download", headers=writer)
    assert downloaded.status_code == 200
    assert downloaded.content.startswith(b"company,report,period,")
    assert b"100.0000" in downloaded.content
    assert downloaded.headers["x-content-sha256"] == artifact.json()["file_sha256"]
    repeated = api_client.post("/reports/generations", headers=headers, json=payload)
    assert repeated.status_code == 201 and repeated.json()["id"] == run["id"]
    history = api_client.get("/reports/generations", headers=owner)
    assert history.status_code == 200
    assert [item["id"] for item in history.json()].count(run["id"]) == 1
    changed = api_client.post(
        "/reports/generations", headers=headers, json={**payload, "format": "PDF"}
    )
    assert changed.status_code == 409


def test_report_catalog_and_run_are_company_scoped(api_client, auth_service):
    owner, writer, period, _, _, _ = closed_company(api_client, auth_service)
    catalog = api_client.get("/reports/catalog", headers=writer)
    assert catalog.status_code == 200
    assert {row["code"] for row in catalog.json()} == {
        "BALANCE_SHEET",
        "INCOME_STATEMENT",
        "FINANCIAL_METRICS",
        "CASH_FLOW",
        "CHANGES_IN_EQUITY",
        "NOTES",
    }
    assert all(row["can_generate"] for row in catalog.json())
    generated = api_client.post(
        "/financial-statements/generate",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json={"period_id": period},
    )
    assert generated.status_code == 201
    period_rows = api_client.get("/accounting/periods", headers=writer).json()
    row = next(item for item in period_rows if item["id"] == period)
    result = api_client.post(
        "/reports/generations",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json={
            "report_code": "INCOME_STATEMENT",
            "format": "XLSX",
            "period_id": period,
            "period_preset": "CUSTOM",
            "date_from": row["start_date"],
            "date_to": row["end_date"],
        },
    )
    assert result.status_code == 201, result.text
    foreign_account = api_client.post(
        "/auth/register",
        json={"email": "foreign-report@example.com", "password": "StrongPassword!2026"},
    )
    assert foreign_account.status_code == 201
    foreign_owner = {
        **owner,
        "Authorization": "Bearer " + foreign_account.json()["access_token"],
    }
    foreign = api_client.get(f"/reports/generations/{result.json()['id']}", headers=foreign_owner)
    assert foreign.status_code == 404
    assert api_client.get("/reports/generations", headers=foreign_owner).status_code == 404
    assert (
        api_client.get(
            f"/reports/generations/{result.json()['id']}/download", headers=foreign_owner
        ).status_code
        == 404
    )
    assert (
        api_client.get(f"/reports/generations/{result.json()['id']}", headers=owner).status_code
        == 200
    )
    downloaded = api_client.get(
        f"/reports/generations/{result.json()['id']}/download", headers=owner
    )
    assert downloaded.status_code == 200
    workbook = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Summary", "Dataset"]
    assert workbook["Summary"]["B2"].value == "INCOME_STATEMENT"
    assert workbook["Dataset"]["A1"].value == "line_code"


def test_metrics_pdf_and_failed_storage_retry(api_client, auth_service, monkeypatch):
    _, writer, period, _, _, _ = closed_company(api_client, auth_service)
    generated = api_client.post(
        "/financial-statements/generate",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json={"period_id": period},
    )
    assert generated.status_code == 201
    row = next(
        item
        for item in api_client.get("/accounting/periods", headers=writer).json()
        if item["id"] == period
    )
    payload = {
        "report_code": "FINANCIAL_METRICS",
        "format": "PDF",
        "period_id": period,
        "period_preset": "CUSTOM",
        "date_from": row["start_date"],
        "date_to": row["end_date"],
    }
    storage = api_client.app.state.services.report_center.storage
    original_put = storage.put

    def unavailable(*_args):
        raise StorageUnavailable()

    monkeypatch.setattr(storage, "put", unavailable)
    failed = api_client.post(
        "/reports/generations",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json=payload,
    )
    assert failed.status_code == 201, failed.text
    assert failed.json()["status"] == "FAILED"
    assert failed.json()["error_code"] == "STORAGE_UNAVAILABLE"
    assert (
        api_client.get(
            f"/reports/generations/{failed.json()['id']}/download", headers=writer
        ).status_code
        == 404
    )
    monkeypatch.setattr(storage, "put", original_put)
    retried = api_client.post(
        f"/reports/generations/{failed.json()['id']}/retry",
        headers={**writer, "Idempotency-Key": str(uuid4())},
    )
    assert retried.status_code == 201, retried.text
    assert retried.json()["retry_of_run_id"] == failed.json()["id"]
    assert retried.json()["source_digest"] == failed.json()["source_digest"]
    completed = api_client.post(
        f"/reports/generations/{retried.json()['id']}/execute", headers=writer
    )
    assert completed.status_code == 200, completed.text
    assert completed.json()["status"] == "COMPLETED"
    downloaded = api_client.get(
        f"/reports/generations/{completed.json()['id']}/download", headers=writer
    )
    assert downloaded.status_code == 200
    assert downloaded.content.startswith(b"%PDF-")
