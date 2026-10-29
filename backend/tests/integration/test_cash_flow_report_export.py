"""Report Center exports only a pinned, reconciled cash flow snapshot."""

from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.reporting.cash_flow.models import CashFlowClassificationModel
from tests.integration.test_cash_flow_snapshot import ready_snapshot
from tests.integration.test_closing_workflow import request


@pytest.mark.parametrize("format", ["CSV", "XLSX", "PDF"])
def test_cash_flow_report_uses_frozen_snapshot(api_client, auth_service, format):
    owner, writer, payload, _, connection = ready_snapshot(api_client, auth_service)
    snapshot_response = request(api_client, writer, "/cash-flow/generate", payload)
    assert snapshot_response.status_code == 201, snapshot_response.text
    snapshot = snapshot_response.json()
    period = next(
        item
        for item in api_client.get("/accounting/periods", headers=writer).json()
        if item["id"] == payload["period_id"]
    )
    catalog = api_client.get("/reports/catalog", headers=writer)
    assert catalog.status_code == 200
    assert (
        next(item for item in catalog.json() if item["code"] == "CASH_FLOW")["period_mode"]
        == "RANGE"
    )
    run_payload = {
        "report_code": "CASH_FLOW",
        "format": format,
        "period_id": payload["period_id"],
        "period_preset": "CUSTOM",
        "date_from": period["start_date"],
        "date_to": period["end_date"],
    }
    headers = {**writer, "Idempotency-Key": str(uuid4())}
    created = api_client.post("/reports/generations", headers=headers, json=run_payload)
    assert created.status_code == 201, created.text
    run = created.json()
    assert run["status"] == "COMPLETED"
    assert run["source_digest"] == snapshot["snapshot_digest"]
    repeated = api_client.post("/reports/generations", headers=headers, json=run_payload)
    assert repeated.status_code == 201 and repeated.json()["id"] == run["id"]
    downloaded = api_client.get(f"/reports/generations/{run['id']}/download", headers=owner)
    assert downloaded.status_code == 200
    if format == "CSV":
        content = downloaded.content.decode("utf-8")
        assert "INTERNAL_VALIDATION_ONLY" in content
        assert "CLOSING_CASH" in content
        assert snapshot["closing_cash"] in content
    elif format == "XLSX":
        book = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=True)
        assert book["Dataset"]["A1"].value == "line_code"
        assert any(row[0] == "CLOSING_CASH" for row in book["Dataset"].iter_rows(values_only=True))
    else:
        assert downloaded.content.startswith(b"%PDF-")
    with Session(connection, join_transaction_mode="create_savepoint") as session:
        classification = session.scalar(
            select(CashFlowClassificationModel).where(
                CashFlowClassificationModel.company_id == UUID(writer["X-Company-ID"])
            )
        )
        assert classification is not None
        classification.version += 1
        session.commit()
    stale = api_client.post(
        "/reports/generations",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json=run_payload,
    )
    assert stale.status_code == 409
    assert (
        api_client.get(f"/reports/generations/{run['id']}/download", headers=owner).content
        == downloaded.content
    )


def test_cash_flow_export_requires_snapshot(api_client, auth_service):
    _, writer, payload, _, _ = ready_snapshot(api_client, auth_service)
    period = next(
        item
        for item in api_client.get("/accounting/periods", headers=writer).json()
        if item["id"] == payload["period_id"]
    )
    response = api_client.post(
        "/reports/generations",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json={
            "report_code": "CASH_FLOW",
            "format": "CSV",
            "period_id": payload["period_id"],
            "period_preset": "CUSTOM",
            "date_from": period["start_date"],
            "date_to": period["end_date"],
        },
    )
    assert response.status_code == 404
