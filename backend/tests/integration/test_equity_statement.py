"""CLOSED BS/IS to frozen SCE, provenance and company isolation."""

from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.reporting.disclosures.models import DisclosureNarrativeModel, DisclosureValueModel
from app.reporting.equity.models import EquityStatementModel
from tests.integration.test_closing_workflow import request, reviewer
from tests.integration.test_finance_flow import post_journal
from tests.integration.test_journal_flow import prepare


def close(client, owner, writer, review, period_id):
    run = request(client, writer, "/closing-runs", {"accounting_period_id": period_id})
    assert run.status_code == 201, run.text
    approval = request(
        client, review, f"/closing-runs/{run.json()['id']}/submit", {"reason": "월마감"}
    )
    assert approval.status_code == 200, approval.text
    assert (
        request(
            client, owner, f"/closing-approvals/{approval.json()['id']}/approve", {"reason": "확인"}
        ).status_code
        == 200
    )
    result = request(
        client,
        owner,
        f"/closing-runs/{run.json()['id']}/close",
        {"approval_id": approval.json()["id"]},
    )
    assert result.status_code == 200, result.text


@pytest.mark.parametrize("physical_transfer", [False, True])
def test_equity_rollforward_uses_prior_bs_and_current_is(
    api_client, auth_service, physical_transfer
):
    owner, writer, payload, connection = prepare(api_client, auth_service)
    review = reviewer(api_client, owner, connection)
    periods = api_client.get("/closing/periods", headers=writer).json()
    january = next(p for p in periods if p["period_no"] == 1)
    february = next(p for p in periods if p["period_no"] == 2)
    accounts = {
        a["account_name"]: a["id"] for a in api_client.get("/accounts", headers=writer).json()
    }
    capital = post_journal(
        api_client,
        owner,
        writer,
        {
            "accounting_period_id": january["id"],
            "entry_date": "2026-01-02",
            "source_type": "MANUAL",
        },
        [
            {"account_id": accounts["현금및현금성자산"], "debit_amount": "200.0000"},
            {"account_id": accounts["자본금"], "credit_amount": "200.0000"},
        ],
    )
    post_journal(
        api_client,
        owner,
        writer,
        {
            "accounting_period_id": january["id"],
            "entry_date": "2026-01-03",
            "source_type": "MANUAL",
        },
        [
            {"account_id": accounts["현금및현금성자산"], "debit_amount": "100.0000"},
            {"account_id": accounts["매출"], "credit_amount": "100.0000"},
        ],
    )
    close(api_client, owner, writer, review, january["id"])
    assert (
        request(
            api_client, writer, "/financial-statements/generate", {"period_id": january["id"]}
        ).status_code
        == 201
    )
    post_journal(
        api_client,
        owner,
        writer,
        {
            "accounting_period_id": february["id"],
            "entry_date": "2026-02-03",
            "source_type": "MANUAL",
        },
        [
            {"account_id": accounts["현금및현금성자산"], "debit_amount": "50.0000"},
            {"account_id": accounts["매출"], "credit_amount": "50.0000"},
        ],
    )
    contribution = post_journal(
        api_client,
        owner,
        writer,
        {
            "accounting_period_id": february["id"],
            "entry_date": "2026-02-04",
            "source_type": "MANUAL",
        },
        [
            {"account_id": accounts["현금및현금성자산"], "debit_amount": "25.0000"},
            {"account_id": accounts["자본금"], "credit_amount": "25.0000"},
        ],
    )
    transfer = None
    if physical_transfer:
        transfer = post_journal(
            api_client,
            owner,
            writer,
            {
                "accounting_period_id": february["id"],
                "entry_date": "2026-02-04",
                "source_type": "MANUAL",
            },
            [
                {"account_id": accounts["매출"], "debit_amount": "100.0000"},
                {"account_id": accounts["미처분이익잉여금"], "credit_amount": "100.0000"},
            ],
        )
    dividend = post_journal(
        api_client,
        owner,
        writer,
        {
            "accounting_period_id": february["id"],
            "entry_date": "2026-02-05",
            "source_type": "MANUAL",
        },
        [
            {"account_id": accounts["미처분이익잉여금"], "debit_amount": "10.0000"},
            {"account_id": accounts["미지급배당금"], "credit_amount": "10.0000"},
        ],
    )
    close(api_client, owner, writer, review, february["id"])
    assert (
        request(
            api_client, writer, "/financial-statements/generate", {"period_id": february["id"]}
        ).status_code
        == 201
    )
    retained = api_client.post(
        "/equity/components",
        headers=writer,
        json={
            "component_code": "RETAINED",
            "name": "이익잉여금",
            "component_type": "RETAINED_EARNINGS",
            "display_order": 2,
        },
    )
    assert retained.status_code == 201, retained.text
    contributed = api_client.post(
        "/equity/components",
        headers=writer,
        json={
            "component_code": "CAPITAL",
            "name": "납입자본",
            "component_type": "CONTRIBUTED_CAPITAL",
            "display_order": 1,
        },
    )
    assert contributed.status_code == 201, contributed.text
    blocked = api_client.get(
        "/equity/reconciliation", headers=owner, params={"period_id": february["id"]}
    )
    assert blocked.status_code == 200 and blocked.json()["status"] == "BLOCKED"
    assert "EQUITY_MAPPING_INCOMPLETE" in blocked.json()["issues"]
    mapped = api_client.post(
        "/equity/mappings",
        headers=writer,
        json={
            "component_id": contributed.json()["id"],
            "account_id": accounts["자본금"],
            "effective_from": "2026-01-01",
        },
    )
    assert mapped.status_code == 201, mapped.text
    retained_mapping = api_client.post(
        "/equity/mappings",
        headers=writer,
        json={
            "component_id": retained.json()["id"],
            "account_id": accounts["미처분이익잉여금"],
            "effective_from": "2026-01-01",
        },
    )
    assert retained_mapping.status_code == 201, retained_mapping.text
    unclassified = api_client.get(
        "/equity/reconciliation", headers=owner, params={"period_id": february["id"]}
    )
    assert "UNCLASSIFIED_EQUITY_SOURCE" in unclassified.json()["issues"]
    declared = api_client.post(
        "/equity/dividends",
        headers=writer,
        json={
            "resolution_date": "2026-02-04",
            "payable_date": "2026-03-10",
            "total_amount": "10.0000",
            "source_journal_id": dividend["id"],
        },
    )
    assert declared.status_code == 201, declared.text
    assert (
        api_client.get(f"/equity/dividends/{declared.json()['id']}/payments", headers=owner).json()
        == []
    )
    for journal, account_id, kind in [
        (contribution, accounts["자본금"], "CAPITAL_CONTRIBUTION"),
        (dividend, accounts["미처분이익잉여금"], "DIVIDEND"),
    ]:
        line_id = next(line["id"] for line in journal["lines"] if line["account_id"] == account_id)
        classified = api_client.post(
            "/equity/events", headers=writer, json={"journal_line_id": line_id, "event_type": kind}
        )
        assert classified.status_code == 201, classified.text
    if transfer:
        transfer_line = next(
            line["id"]
            for line in transfer["lines"]
            if line["account_id"] == accounts["미처분이익잉여금"]
        )
        classified = api_client.post(
            "/equity/events",
            headers=writer,
            json={"journal_line_id": transfer_line, "event_type": "PROFIT_TRANSFER"},
        )
        assert classified.status_code == 201, classified.text
    ready = api_client.get(
        "/equity/reconciliation", headers=owner, params={"period_id": february["id"]}
    )
    assert ready.status_code == 200, ready.text
    if transfer:
        assert ready.json()["status"] == "BLOCKED", ready.json()
        assert "PROFIT_TRANSFER_IS_MISMATCH" in ready.json()["issues"]
        assert Decimal(ready.json()["income_statement_profit"]) == Decimal("-50")
        assert (
            request(
                api_client, writer, "/equity/statements/generate", {"period_id": february["id"]}
            ).status_code
            == 422
        )
        return
    assert ready.json()["status"] == "READY", ready.json()
    assert Decimal(ready.json()["balance_sheet_equity"]) == Decimal("365")
    assert Decimal(ready.json()["income_statement_profit"]) == Decimal("50")
    assert Decimal(ready.json()["rollforward_equity"]) == Decimal("365")
    key = str(uuid4())
    generated = request(
        api_client, writer, "/equity/statements/generate", {"period_id": february["id"]}, key
    )
    assert generated.status_code == 201, generated.text
    replay = request(
        api_client, writer, "/equity/statements/generate", {"period_id": february["id"]}, key
    )
    assert replay.status_code == 201 and replay.json()["id"] == generated.json()["id"]
    detail = api_client.get(f"/equity/statements/{generated.json()['id']}", headers=owner)
    assert detail.status_code == 200
    lines = {(line["component_id"], line["movement_type"]): line for line in detail.json()["lines"]}
    assert Decimal(lines[(retained.json()["id"], "OPENING")]["amount"]) == Decimal("100")
    assert Decimal(lines[(retained.json()["id"], "PROFIT_OR_LOSS")]["amount"]) == Decimal("50")
    assert Decimal(lines[(contributed.json()["id"], "CLOSING")]["amount"]) == Decimal("225")
    assert Decimal(lines[(retained.json()["id"], "OWNER_TRANSACTIONS")]["amount"]) == Decimal("-10")
    opening_line = lines[(contributed.json()["id"], "OPENING")]
    sources = api_client.get(
        f"/equity/statements/{generated.json()['id']}/lines/{opening_line['id']}/sources",
        headers=owner,
    )
    assert sources.status_code == 200 and sources.json()[0]["journal_id"] == capital["id"]
    assert api_client.get(
        f"/equity/statements/{generated.json()['id']}",
        headers={**owner, "X-Company-ID": str(uuid4())},
    ).status_code in {403, 404}
    with Session(connection, join_transaction_mode="create_savepoint") as session:
        with pytest.raises(IntegrityError):
            session.execute(
                update(EquityStatementModel)
                .where(EquityStatementModel.id == UUID(generated.json()["id"]))
                .values(status="RECONCILED")
            )
        session.rollback()

    march = next(p for p in periods if p["period_no"] == 3)
    first_payment = post_journal(
        api_client,
        owner,
        writer,
        {"accounting_period_id": march["id"], "entry_date": "2026-03-04", "source_type": "MANUAL"},
        [
            {"account_id": accounts["미지급배당금"], "debit_amount": "5.0000"},
            {"account_id": accounts["현금및현금성자산"], "credit_amount": "5.0000"},
        ],
    )
    excess_payment = post_journal(
        api_client,
        owner,
        writer,
        {"accounting_period_id": march["id"], "entry_date": "2026-03-05", "source_type": "MANUAL"},
        [
            {"account_id": accounts["미지급배당금"], "debit_amount": "6.0000"},
            {"account_id": accounts["현금및현금성자산"], "credit_amount": "6.0000"},
        ],
    )
    close(api_client, owner, writer, review, march["id"])
    assert (
        request(
            api_client, writer, "/financial-statements/generate", {"period_id": march["id"]}
        ).status_code
        == 201
    )
    assert (
        api_client.post(
            "/cash-flow/cash-scope",
            headers=writer,
            json={
                "account_id": accounts["현금및현금성자산"],
                "scope_type": "CASH",
                "effective_from": "2026-01-01",
            },
        ).status_code
        == 201
    )
    assert (
        api_client.post(
            "/cash-flow/classifications",
            headers=writer,
            json={
                "account_id": accounts["미지급배당금"],
                "activity_type": "FINANCING",
                "line_code": "DIVIDENDS_PAID",
                "effective_from": "2026-03-01",
            },
        ).status_code
        == 201
    )
    assert (
        api_client.post(
            "/cash-flow/adjustment-rules",
            headers=writer,
            json={
                "account_id": accounts["미지급배당금"],
                "adjustment_code": "EXCLUDED_BALANCE",
                "source_mapping_type": "ACCOUNT_LEDGER",
                "related_account_id": None,
                "effective_from": "2026-03-01",
                "policy_code": "TAX_INTEREST_OPERATING",
                "policy_version": 1,
            },
        ).status_code
        == 201
    )
    bridge = api_client.get(
        "/cash-flow/indirect-bridge", headers=writer, params={"period_id": march["id"]}
    )
    assert bridge.status_code == 200 and bridge.json()["status"] == "READY", bridge.text
    cf_payload = {
        key: bridge.json()[key]
        for key in (
            "period_id",
            "statement_set_id",
            "close_version",
            "source_digest",
            "mapping_digest",
        )
    }
    cash_flow = request(api_client, writer, "/cash-flow/generate", cf_payload)
    assert cash_flow.status_code == 201, cash_flow.text
    events = api_client.get(
        f"/cash-flow/{cash_flow.json()['id']}/drilldown", headers=writer
    ).json()["events"]
    event_ids = {event["journal_id"]: event["id"] for event in events}
    assert set(event_ids) == {first_payment["id"], excess_payment["id"]}
    payment_url = f"/equity/dividends/{declared.json()['id']}/payments"
    paid = api_client.post(
        payment_url, headers=writer, json={"cash_flow_event_id": event_ids[first_payment["id"]]}
    )
    assert paid.status_code == 201, paid.text
    assert Decimal(paid.json()["amount"]) == Decimal("5")
    replayed_payment = api_client.post(
        payment_url, headers=writer, json={"cash_flow_event_id": event_ids[first_payment["id"]]}
    )
    assert (
        replayed_payment.status_code == 201 and replayed_payment.json()["id"] == paid.json()["id"]
    )
    blocked_payment = api_client.post(
        payment_url, headers=writer, json={"cash_flow_event_id": event_ids[excess_payment["id"]]}
    )
    assert blocked_payment.status_code == 409, blocked_payment.text
    assert len(api_client.get(payment_url, headers=writer).json()) == 1
    assert api_client.post(
        payment_url,
        headers={**writer, "X-Company-ID": str(uuid4())},
        json={"cash_flow_event_id": event_ids[first_payment["id"]]},
    ).status_code in {403, 404}

    march_equity_ready = api_client.get(
        "/equity/reconciliation", headers=writer, params={"period_id": march["id"]}
    )
    assert march_equity_ready.status_code == 200, march_equity_ready.text
    assert march_equity_ready.json()["status"] == "READY", march_equity_ready.json()
    march_equity = request(
        api_client, writer, "/equity/statements/generate", {"period_id": march["id"]}
    )
    assert march_equity.status_code == 201, march_equity.text

    notes_ready = api_client.get(
        "/disclosures/readiness", headers=owner, params={"period_id": march["id"]}
    )
    assert notes_ready.status_code == 200, notes_ready.text
    assert notes_ready.json()["status"] == "READY", notes_ready.json()
    template = api_client.get("/disclosures/templates", headers=owner)
    assert template.status_code == 200, template.text
    sections = template.json()[0]["sections"]
    notes_key = str(uuid4())
    notes = request(
        api_client,
        owner,
        "/disclosures/instances/generate",
        {"period_id": march["id"]},
        notes_key,
    )
    assert notes.status_code == 201, notes.text
    notes_instance = notes.json()
    assert notes_instance["status"] == "DRAFT"
    assert notes_instance["reconciliation_status"] == "RECONCILED"
    assert len(notes_instance["values"]) > len(sections)
    account_facts = [
        value
        for value in notes_instance["values"]
        if value.get("account_id") and Decimal(value["decimal_value"]) != 0
    ]
    assert account_facts
    account_fact = account_facts[0]
    account_source = api_client.get(
        f"/disclosures/instances/{notes_instance['id']}/values/{account_fact['id']}/source",
        headers=writer,
    )
    assert account_source.status_code == 200, account_source.text
    assert account_source.json()["account_id"] == account_fact["account_id"]
    assert account_source.json()["journal_line_ids"]
    assert account_source.json()["amount"] == account_fact["decimal_value"]
    replay = request(
        api_client,
        owner,
        "/disclosures/instances/generate",
        {"period_id": march["id"]},
        notes_key,
    )
    assert replay.status_code == 201 and replay.json()["id"] == notes_instance["id"]
    conflicting_replay = request(
        api_client,
        owner,
        "/disclosures/instances/generate",
        {"period_id": february["id"]},
        notes_key,
    )
    assert conflicting_replay.status_code == 409, conflicting_replay.text
    source_value = next(
        value for value in notes_instance["values"] if value["source_type"] == "BALANCE_SHEET"
    )
    source = api_client.get(
        f"/disclosures/instances/{notes_instance['id']}/values/{source_value['id']}/source",
        headers=owner,
    )
    assert source.status_code == 200, source.text
    assert source.json()["source_id"] == source_value["source_id"]
    assert source.json()["source_path"].startswith("/financial-statements/sets/")
    assert api_client.get(
        f"/disclosures/instances/{notes_instance['id']}/values/{source_value['id']}/source",
        headers={**writer, "X-Company-ID": str(uuid4())},
    ).status_code in {403, 404}

    instance_version = notes_instance["version"]
    for section in sections:
        saved = api_client.put(
            f"/disclosures/instances/{notes_instance['id']}/narratives",
            headers=owner,
            json={
                "section_id": section["id"],
                "text_value": f"{section['title']}의 검증된 Fact를 검토했습니다.",
                "expected_version": instance_version,
            },
        )
        assert saved.status_code == 200, saved.text
        instance_version = saved.json()["instance_version"]
    stale_review = api_client.post(
        f"/disclosures/instances/{notes_instance['id']}/request-review",
        headers=owner,
        json={"expected_version": notes_instance["version"]},
    )
    assert stale_review.status_code == 409, stale_review.text
    requested = api_client.post(
        f"/disclosures/instances/{notes_instance['id']}/request-review",
        headers=owner,
        json={"expected_version": instance_version},
    )
    assert requested.status_code == 200, requested.text
    assert requested.json()["status"] == "REVIEW_REQUESTED"
    assert (
        api_client.post(
            f"/disclosures/instances/{notes_instance['id']}/approve",
            headers=owner,
            json={"expected_version": instance_version},
        ).status_code
        == 409
    )
    assert (
        api_client.post(
            f"/disclosures/instances/{notes_instance['id']}/approve",
            headers=writer,
            json={"expected_version": requested.json()["version"]},
        ).status_code
        == 403
    )
    rejected = api_client.post(
        f"/disclosures/instances/{notes_instance['id']}/reject",
        headers=review,
        json={"expected_version": requested.json()["version"], "reason": "표현 보완"},
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["status"] == "REJECTED"
    revised_narrative = api_client.put(
        f"/disclosures/instances/{notes_instance['id']}/narratives",
        headers=owner,
        json={
            "section_id": sections[0]["id"],
            "text_value": "검토 의견을 반영해 원천을 다시 확인했습니다.",
            "expected_version": rejected.json()["version"],
        },
    )
    assert revised_narrative.status_code == 200, revised_narrative.text
    requested_again = api_client.post(
        f"/disclosures/instances/{notes_instance['id']}/request-review",
        headers=owner,
        json={"expected_version": revised_narrative.json()["instance_version"]},
    )
    assert requested_again.status_code == 200, requested_again.text
    approved = api_client.post(
        f"/disclosures/instances/{notes_instance['id']}/approve",
        headers=review,
        json={"expected_version": requested_again.json()["version"]},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "APPROVED"
    assert approved.json()["approved_by"] != approved.json()["created_by"]
    with Session(connection, join_transaction_mode="create_savepoint") as session:
        with pytest.raises(IntegrityError):
            session.execute(
                update(DisclosureValueModel)
                .where(DisclosureValueModel.id == UUID(source_value["id"]))
                .values(decimal_value=Decimal("0"))
            )
        session.rollback()
        with pytest.raises(IntegrityError):
            session.execute(
                update(DisclosureNarrativeModel)
                .where(
                    DisclosureNarrativeModel.instance_id == UUID(notes_instance["id"]),
                    DisclosureNarrativeModel.section_id == UUID(sections[0]["id"]),
                )
                .values(text_value="변경 불가")
            )
        session.rollback()

    for report_format in ("CSV", "XLSX", "PDF"):
        report = api_client.post(
            "/reports/generations",
            headers={**writer, "Idempotency-Key": str(uuid4())},
            json={
                "report_code": "NOTES",
                "format": report_format,
                "period_id": march["id"],
                "period_preset": "CUSTOM",
                "date_from": march["start_date"],
                "date_to": march["end_date"],
            },
        )
        assert report.status_code == 201, report.text
        artifact = api_client.get(
            f"/reports/generations/{report.json()['id']}/artifact", headers=writer
        )
        assert artifact.status_code == 200, artifact.text
        download = api_client.get(
            f"/reports/generations/{report.json()['id']}/download", headers=writer
        )
        assert download.status_code == 200
        assert download.headers["x-content-sha256"] == artifact.json()["file_sha256"]
        assert (
            notes_instance["source_digest"] in download.content.decode("utf-8", errors="ignore")
            if report_format == "CSV"
            else True
        )
        if report_format == "CSV":
            assert "statement_line_id" in download.content.decode("utf-8")
            assert "account_id" in download.content.decode("utf-8")

    april = next(p for p in periods if p["period_no"] == 4)
    invalid_declaration = post_journal(
        api_client,
        owner,
        writer,
        {"accounting_period_id": april["id"], "entry_date": "2026-04-02", "source_type": "MANUAL"},
        [
            {"account_id": accounts["미처분이익잉여금"], "debit_amount": "4.0000"},
            {"account_id": accounts["현금및현금성자산"], "credit_amount": "4.0000"},
        ],
    )
    assert (
        api_client.post(
            "/equity/dividends",
            headers=writer,
            json={
                "resolution_date": "2026-04-02",
                "total_amount": "4.0000",
                "source_journal_id": invalid_declaration["id"],
            },
        ).status_code
        == 409
    )
    catalog = api_client.get("/reports/catalog", headers=owner).json()
    assert (
        next(item for item in catalog if item["code"] == "CHANGES_IN_EQUITY")["period_mode"]
        == "RANGE"
    )
    csv_run_id = None
    for format in ("CSV", "XLSX", "PDF"):
        report = api_client.post(
            "/reports/generations",
            headers={**writer, "Idempotency-Key": str(uuid4())},
            json={
                "report_code": "CHANGES_IN_EQUITY",
                "format": format,
                "period_id": february["id"],
                "period_preset": "CUSTOM",
                "date_from": february["start_date"],
                "date_to": february["end_date"],
            },
        )
        assert report.status_code == 201, report.text
        assert report.json()["status"] == "COMPLETED", report.text
        if format == "CSV":
            csv_run_id = report.json()["id"]
        assert report.json()["source_digest"] == generated.json()["snapshot_digest"]
        download = api_client.get(
            f"/reports/generations/{report.json()['id']}/download",
            headers=owner,
        )
        assert download.status_code == 200
        assert download.headers["X-Content-SHA256"] == sha256(download.content).hexdigest()
        if format == "CSV":
            content = download.content.decode("utf-8")
            assert "INTERNAL_VALIDATION_ONLY" in content
            assert "CLOSING" in content and "225.0000" in content and "140.0000" in content
        elif format == "XLSX":
            book = load_workbook(BytesIO(download.content), read_only=True, data_only=True)
            assert book["Dataset"]["A1"].value == "movement_type"
            assert any(row[0] == "CLOSING" for row in book["Dataset"].iter_rows(values_only=True))
        else:
            assert download.content.startswith(b"%PDF-")
        assert api_client.get(
            f"/reports/generations/{report.json()['id']}/download",
            headers={**owner, "X-Company-ID": str(uuid4())},
        ).status_code in {403, 404}
    assert csv_run_id is not None
    history = api_client.get("/reports/generations", headers=owner)
    assert history.status_code == 200
    assert csv_run_id in {run["id"] for run in history.json()}
    assert (
        api_client.get(f"/reports/generations/{csv_run_id}/artifact", headers=owner).status_code
        == 200
    )
    changed = api_client.post(
        "/equity/components",
        headers=writer,
        json={
            "component_code": "OTHER",
            "name": "기타자본",
            "component_type": "OTHER_EQUITY",
            "display_order": 3,
        },
    )
    assert changed.status_code == 201, changed.text
    stale = api_client.post(
        "/reports/generations",
        headers={**writer, "Idempotency-Key": str(uuid4())},
        json={
            "report_code": "CHANGES_IN_EQUITY",
            "format": "CSV",
            "period_id": february["id"],
            "period_preset": "CUSTOM",
            "date_from": february["start_date"],
            "date_to": february["end_date"],
        },
    )
    assert stale.status_code == 409
    assert (
        api_client.get(f"/reports/generations/{csv_run_id}/download", headers=owner).status_code
        == 200
    )
