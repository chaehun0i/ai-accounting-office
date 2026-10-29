"""Report exports preserve a typed dataset across all supported formats."""

from io import BytesIO

from openpyxl import load_workbook

from app.reporting.center.render import ReportDocument, render


def test_csv_xlsx_pdf_share_metadata_and_exact_decimal_text():
    document = ReportDocument(
        metadata=(("company", "샘플회사"), ("snapshot_version", "2")),
        columns=("line_code", "label", "amount"),
        rows=(("ASSETS", "자산총계", "123456789012345.6789"),),
    )
    csv_content, csv_type, csv_extension = render(document, "CSV")
    assert csv_content.decode("utf-8") == (
        "company,snapshot_version,line_code,label,amount\n"
        "샘플회사,2,ASSETS,자산총계,123456789012345.6789\n"
    )
    assert csv_type.startswith("text/csv") and csv_extension == "csv"

    xlsx_content, _, xlsx_extension = render(document, "XLSX")
    workbook = load_workbook(BytesIO(xlsx_content), read_only=True, data_only=True)
    assert xlsx_extension == "xlsx"
    assert workbook.sheetnames == ["Summary", "Dataset"]
    assert workbook["Summary"]["B1"].value == "샘플회사"
    assert workbook["Dataset"]["C2"].value == "123456789012345.6789"

    pdf_content, pdf_type, pdf_extension = render(document, "PDF")
    assert pdf_content.startswith(b"%PDF-")
    assert pdf_type == "application/pdf" and pdf_extension == "pdf"


def test_non_cash_disclosure_is_separate_from_cash_flow_totals():
    document = ReportDocument(
        metadata=(("report", "CASH_FLOW"),),
        columns=("line_code", "amount", "source_type"),
        rows=(("NET_OPERATING", "100.0000", "TOTAL"),),
        sections=(
            (
                "NonCashDisclosure",
                ("source_type", "amount", "journal_id", "evidence_id"),
                (("LEASE_IN_KIND", "25.0000", "journal-1", "evidence-1"),),
            ),
        ),
    )
    csv_content, _, _ = render(document, "CSV")
    text = csv_content.decode("utf-8")
    assert "CASH_FLOW,NET_OPERATING,100.0000,TOTAL\n" in text
    assert "\nNonCashDisclosure\nsource_type,amount,journal_id,evidence_id\n" in text
    assert "LEASE_IN_KIND,25.0000,journal-1,evidence-1\n" in text
    xlsx_content, _, _ = render(document, "XLSX")
    workbook = load_workbook(BytesIO(xlsx_content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Summary", "Dataset", "NonCashDisclosure"]
    assert workbook["Dataset"]["B2"].value == "100.0000"
    assert workbook["NonCashDisclosure"]["B2"].value == "25.0000"
    pdf_content, _, _ = render(document, "PDF")
    assert pdf_content.startswith(b"%PDF-")
